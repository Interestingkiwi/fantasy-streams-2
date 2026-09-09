"""
Routes used by Draft Prep
Author - Jason Druckenmiller
Created - 7/3/2026
Updated - 9/8/2026
"""

import io
import re

from flask import Blueprint, render_template, jsonify, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from db import engine, text
from ranking_utils import calculate_player_ranks

# Create the Blueprint with a URL prefix
draft_bp = Blueprint('draft', __name__, url_prefix='/draft-prep')

# Shared with the Schedules page so the two cannot drift apart.
from schedule_utils import team_game_counts  # noqa: E402

def get_stat_mappings(conn):
    """Helper function to dynamically classify stats as Skater or Goalie from the DB."""
    query = text("""
        SELECT column_name
        FROM information_schema.columns
        WHERE table_name = 'final_projections'
          AND column_name NOT IN ('id', 'playerId', 'teamAbbrevs', 'positionCode', 'projectedGames', 'fullName', 'productionTrend', 'peripheralTrend', 'projectionSource', 'onNhlRoster', 'eligiblePositions');
    """)
    result = conn.execute(query)
    stats = [row[0] for row in result]

    if not stats:
        return [], []

    count_selects = ", ".join([f'COUNT("{stat}") AS "{stat}"' for stat in stats])

    # Query for Goalies
    g_query = text(f'SELECT {count_selects} FROM final_projections WHERE "positionCode" = \'G\'')
    g_result = dict(conn.execute(g_query).fetchone()._mapping)
    goalie_stats = [stat for stat in stats if g_result[stat] > 0]

    # Query for Skaters
    s_query = text(f'SELECT {count_selects} FROM final_projections WHERE "positionCode" != \'G\'')
    s_result = dict(conn.execute(s_query).fetchone()._mapping)
    skater_stats = [stat for stat in stats if s_result[stat] > 0]

    strict_goalie_stats = [stat for stat in goalie_stats if stat not in skater_stats]

    return skater_stats, strict_goalie_stats


@draft_bp.route('/')
def dashboard():
    """Serves the main draft prep page at /draft-prep/"""
    return render_template('pages/draft-prep.html')

@draft_bp.route('/api/available-stats')
def get_available_stats():
    """Dynamically fetches stat columns and separates them by Skater vs Goalie."""
    try:
        with engine.connect() as conn:
            skater_stats, goalie_stats = get_stat_mappings(conn)

        return jsonify({
            "status": "success",
            "skater_stats": skater_stats,
            "goalie_stats": goalie_stats
        })
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

@draft_bp.route('/api/projections')
def get_projections():
    """Fetches the actual player data and projections."""
    try:
        with engine.connect() as conn:
            query = text("SELECT * FROM final_projections")
            result = conn.execute(query)
            players = [dict(row._mapping) for row in result]

        return jsonify({"status": "success", "data": players})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@draft_bp.route('/api/playoff-schedule', methods=['POST'])
def playoff_schedule():
    """
    Counts each team's games across the selected fantasy playoff weeks.

    Returns games plus light-night games - those falling on a date carrying
    eight games or fewer, the standard fantasy definition, when enough of the
    league is idle to get an extra starter into the lineup. The threshold
    itself lives in schedule_utils; this route must not restate it.
    """
    try:
        weeks = (request.json or {}).get('weeks', [])
        ranges = [(w.get('start'), w.get('end')) for w in weeks
                  if w.get('start') and w.get('end')]

        if not ranges:
            return jsonify({'status': 'success', 'teams': {}, 'average': 0})

        clauses = []
        params = {}
        for index, (start, end) in enumerate(ranges):
            clauses.append(f'("gameDate" BETWEEN :start{index} AND :end{index})')
            params[f'start{index}'] = start
            params[f'end{index}'] = end
        window = ' OR '.join(clauses)

        with engine.connect() as conn:
            rows = conn.execute(text(f'''
                SELECT "gameDate", "homeTeam", "awayTeam"
                FROM nhl_schedule
                WHERE {window}
            '''), params).fetchall()

        teams = team_game_counts(rows)

        counts = [entry['games'] for entry in teams.values()]
        average = round(sum(counts) / len(counts), 1) if counts else 0

        return jsonify({
            'status': 'success',
            'teams': teams,
            'average': average,
            'max': max(counts) if counts else 0,
            'min': min(counts) if counts else 0,
        })
    except Exception as e:
        return jsonify({'status': 'error', 'message': str(e)}), 500


# --- Export ---------------------------------------------------------------
# The page sends the values and the cell colours it is already showing rather
# than a description of the settings behind them. The heatmap scale, the
# Favourite/Sleeper/DND tags and the polarity of a category all live in the
# browser (localStorage), and re-deriving them here would be a second
# implementation of the same maths that could only ever drift from the first.
# So this route's whole job is turning a grid of values and fills into a
# workbook, and it is the only place that knows anything about xlsx.

MAX_EXPORT_ROWS = 20000
MAX_EXPORT_COLUMNS = 200

# Excel rejects these in a sheet name, and caps the name at 31 characters
INVALID_SHEET_CHARS = re.compile(r'[\[\]:*?/\\]')

HEADER_FILL = PatternFill('solid', fgColor='1F2937')
HEADER_FONT = Font(bold=True, color='FFFFFF')
THIN_EDGE = Side(style='thin', color='D1D5DB')
CELL_BORDER = Border(left=THIN_EDGE, right=THIN_EDGE, top=THIN_EDGE, bottom=THIN_EDGE)


def _sheet_title(name):
    cleaned = INVALID_SHEET_CHARS.sub(' ', str(name or '')).strip()
    return cleaned[:31] or 'Draft List'


def _hex_fill(value):
    """A six-digit RRGGBB from the page, or None. Anything else is dropped."""
    if not isinstance(value, str):
        return None
    candidate = value.strip().lstrip('#').upper()
    if len(candidate) != 6 or not all(c in '0123456789ABCDEF' for c in candidate):
        return None
    return PatternFill('solid', fgColor=candidate)


@draft_bp.route('/api/export', methods=['POST'])
def export_list():
    """Builds an .xlsx of a draft list, colours and all, and returns it."""
    try:
        payload = request.get_json(silent=True) or {}
        columns = payload.get('columns') or []
        rows = payload.get('rows') or []
        name = payload.get('name') or 'Draft List'

        if not columns:
            return jsonify({'status': 'error', 'message': 'No columns to export.'}), 400
        if not rows:
            return jsonify({'status': 'error', 'message': 'No rows to export.'}), 400
        if len(columns) > MAX_EXPORT_COLUMNS or len(rows) > MAX_EXPORT_ROWS:
            return jsonify({'status': 'error', 'message': 'That list is too large to export.'}), 413

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = _sheet_title(name)

        headers = [str(column.get('label', '')) for column in columns]
        # Widths are grown as the rows are written, starting from the header
        widths = [len(header) for header in headers]

        sheet.append(headers)
        for index in range(len(headers)):
            cell = sheet.cell(row=1, column=index + 1)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = CELL_BORDER

        for row_index, row in enumerate(rows, start=2):
            values = (row or {}).get('values') or []
            fills = (row or {}).get('fills') or []

            for column_index in range(len(headers)):
                value = values[column_index] if column_index < len(values) else None
                cell = sheet.cell(row=row_index, column=column_index + 1, value=value)
                cell.border = CELL_BORDER

                fill = _hex_fill(fills[column_index] if column_index < len(fills) else None)
                if fill is not None:
                    cell.fill = fill

                if value is not None:
                    widths[column_index] = max(widths[column_index], len(str(value)))

        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = min(max(width + 2, 8), 40)

        # The header stays put while scrolling, and filters/sorts on its own
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{len(rows) + 1}'

        stream = io.BytesIO()
        workbook.save(stream)
        stream.seek(0)

        return send_file(
            stream,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=f'{_sheet_title(name)}.xlsx',
        )
    except Exception as e:
        return jsonify({'status': 'error', 'message': str(e)}), 500


@draft_bp.route('/api/rank-players', methods=['POST'])
def rank_players():
    try:
        user_settings = request.json
        active_stats = user_settings.get('active_stats', {})

        league_mode = user_settings.get('league_mode', 'categories')

        # Roster shape drives replacement level; absent, ranking falls back to raw value
        num_teams = user_settings.get('num_teams')
        roster_slots = user_settings.get('roster_slots') or {}
        roster_mode = user_settings.get('roster_mode', 'split')

        # roster | projection | balanced - how hard the board leans on scarcity
        rank_mode = user_settings.get('rank_mode', 'roster')

        with engine.connect() as conn:
            skater_stats, goalie_stats = get_stat_mappings(conn)

            query = text("SELECT * FROM final_projections")
            result = conn.execute(query)
            players_data = [dict(row._mapping) for row in result]

            ranked_players = calculate_player_ranks(
                players_data=players_data,
                active_stats=active_stats,
                league_mode=league_mode,
                goalie_stat_keywords=goalie_stats,
                num_teams=num_teams,
                roster_slots=roster_slots,
                roster_mode=roster_mode,
                rank_mode=rank_mode
            )

        return jsonify(ranked_players)

    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500
